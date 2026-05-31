from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import re
import shutil
from datetime import datetime, timedelta
from math import ceil
from pathlib import Path
from typing import Any, Optional

import aiosqlite

from app.models.audit_log import (
    AuditArchiveInfo,
    AuditLog,
    AuditLogPageResult,
    AuditLogQuery,
    AuditLogStatistics,
    AuditLogVerifyResult,
)


class AuditLogService:
    DEFAULT_DB_PATH = Path("data") / "audit_logs.db"
    QUEUE_MAX_SIZE = 10000
    BATCH_FLUSH_INTERVAL = 0.5
    BATCH_MAX_SIZE = 50
    VALID_ORDER_FIELDS = {"timestamp", "user_id", "module", "action_type", "status"}
    ARCHIVE_RETENTION_YEARS = 3
    BACKUP_RETENTION_DAYS = 30

    def __init__(self):
        self._db: Optional[aiosqlite.Connection] = None
        self._db_path: Optional[str] = None
        self._queue: Optional[asyncio.Queue] = None
        self._consumer_task: Optional[asyncio.Task] = None
        self._running = False

    async def initialize(self, db_path: str = None) -> None:
        if db_path is None:
            project_root = Path(__file__).resolve().parent.parent.parent.parent
            db_path = str(project_root / self.DEFAULT_DB_PATH)
        self._db_path = db_path

        db_dir = Path(db_path).parent
        db_dir.mkdir(parents=True, exist_ok=True)

        self._db = await aiosqlite.connect(db_path)
        self._db.row_factory = aiosqlite.Row

        await self._create_tables()

        self._queue = asyncio.Queue(maxsize=self.QUEUE_MAX_SIZE)
        self._running = True
        self.start_queue_consumer()

    async def _create_tables(self) -> None:
        await self._db.executescript("""
            CREATE TABLE IF NOT EXISTS audit_logs (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL DEFAULT 'anonymous',
                username TEXT NOT NULL DEFAULT 'anonymous',
                timestamp TEXT NOT NULL,
                ip_address TEXT NOT NULL DEFAULT '',
                action_type TEXT NOT NULL,
                module TEXT NOT NULL,
                detail TEXT,
                status TEXT NOT NULL DEFAULT 'success',
                client_info TEXT NOT NULL DEFAULT '',
                request_path TEXT NOT NULL DEFAULT '',
                request_method TEXT NOT NULL DEFAULT '',
                response_code INTEGER NOT NULL DEFAULT 200,
                duration_ms REAL NOT NULL DEFAULT 0.0,
                hash TEXT NOT NULL DEFAULT '',
                sensitive INTEGER NOT NULL DEFAULT 0
            );

            CREATE INDEX IF NOT EXISTS idx_audit_logs_timestamp ON audit_logs(timestamp);
            CREATE INDEX IF NOT EXISTS idx_audit_logs_user_id ON audit_logs(user_id);
            CREATE INDEX IF NOT EXISTS idx_audit_logs_module ON audit_logs(module);
            CREATE INDEX IF NOT EXISTS idx_audit_logs_action_type ON audit_logs(action_type);
            CREATE INDEX IF NOT EXISTS idx_audit_logs_status ON audit_logs(status);
            CREATE INDEX IF NOT EXISTS idx_audit_logs_request_path ON audit_logs(request_path);

            CREATE VIRTUAL TABLE IF NOT EXISTS audit_logs_fts USING fts5(
                detail,
                content='audit_logs',
                content_rowid='rowid'
            );

            CREATE TRIGGER IF NOT EXISTS audit_logs_ai AFTER INSERT ON audit_logs BEGIN
                INSERT INTO audit_logs_fts(rowid, detail) VALUES (new.rowid, COALESCE(new.detail, ''));
            END;

            CREATE TRIGGER IF NOT EXISTS audit_logs_ad AFTER DELETE ON audit_logs BEGIN
                INSERT INTO audit_logs_fts(audit_logs_fts, rowid, detail) VALUES('delete', old.rowid, COALESCE(old.detail, ''));
            END;

            CREATE TRIGGER IF NOT EXISTS audit_logs_au AFTER UPDATE ON audit_logs BEGIN
                INSERT INTO audit_logs_fts(audit_logs_fts, rowid, detail) VALUES('delete', old.rowid, COALESCE(old.detail, ''));
                INSERT INTO audit_logs_fts(rowid, detail) VALUES (new.rowid, COALESCE(new.detail, ''));
            END;
        """)
        await self._db.commit()

    async def shutdown(self) -> None:
        await self.stop_queue_consumer()
        if self._queue is not None:
            remaining = []
            while not self._queue.empty():
                try:
                    remaining.append(self._queue.get_nowait())
                except asyncio.QueueEmpty:
                    break
            if remaining:
                await self._write_batch(remaining)
        if self._db is not None:
            await self._db.close()
            self._db = None

    def start_queue_consumer(self) -> None:
        if self._consumer_task is None or self._consumer_task.done():
            self._consumer_task = asyncio.create_task(self._queue_consumer_loop())

    async def stop_queue_consumer(self) -> None:
        self._running = False
        if self._consumer_task is not None and not self._consumer_task.done():
            self._consumer_task.cancel()
            try:
                await self._consumer_task
            except asyncio.CancelledError:
                pass
            self._consumer_task = None

    async def _queue_consumer_loop(self) -> None:
        batch: list[dict] = []
        while self._running:
            try:
                try:
                    entry = await asyncio.wait_for(
                        self._queue.get(), timeout=self.BATCH_FLUSH_INTERVAL
                    )
                    batch.append(entry)
                except asyncio.TimeoutError:
                    pass

                while len(batch) < self.BATCH_MAX_SIZE:
                    try:
                        entry = self._queue.get_nowait()
                        batch.append(entry)
                    except asyncio.QueueEmpty:
                        break

                if batch and (len(batch) >= self.BATCH_MAX_SIZE or not self._running):
                    await self._write_batch(batch)
                    batch = []
                elif batch:
                    await asyncio.sleep(self.BATCH_FLUSH_INTERVAL)
                    await self._write_batch(batch)
                    batch = []

            except asyncio.CancelledError:
                if batch:
                    await self._write_batch(batch)
                break
            except Exception:
                if batch:
                    batch = []
                await asyncio.sleep(1)

    def enqueue_log(self, entry: dict) -> None:
        if self._queue is None:
            return
        try:
            self._queue.put_nowait(entry)
        except asyncio.QueueFull:
            try:
                self._queue.get_nowait()
            except asyncio.QueueEmpty:
                pass
            try:
                self._queue.put_nowait(entry)
            except asyncio.QueueFull:
                pass

    async def _write_batch(self, entries: list[dict]) -> None:
        if not entries or self._db is None:
            return
        await self._db.execute("BEGIN")
        try:
            for entry in entries:
                detail_val = entry.get("detail")
                if detail_val is not None and not isinstance(detail_val, str):
                    detail_val = json.dumps(detail_val, ensure_ascii=False)
                elif detail_val is None:
                    detail_val = None

                sensitive_val = 1 if entry.get("sensitive", False) else 0

                await self._db.execute(
                    """
                    INSERT INTO audit_logs (id, user_id, username, timestamp, ip_address,
                        action_type, module, detail, status, client_info,
                        request_path, request_method, response_code, duration_ms, hash, sensitive)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        entry.get("id", ""),
                        entry.get("user_id", "anonymous"),
                        entry.get("username", "anonymous"),
                        entry.get("timestamp", ""),
                        entry.get("ip_address", ""),
                        entry.get("action_type", ""),
                        entry.get("module", ""),
                        detail_val,
                        entry.get("status", "success"),
                        entry.get("client_info", ""),
                        entry.get("request_path", ""),
                        entry.get("request_method", ""),
                        entry.get("response_code", 200),
                        entry.get("duration_ms", 0.0),
                        entry.get("hash", ""),
                        sensitive_val,
                    ),
                )
            await self._db.commit()
        except Exception:
            await self._db.rollback()
            raise

    @staticmethod
    def compute_hash(entry: dict) -> str:
        fields = [
            "id", "user_id", "username", "timestamp", "ip_address",
            "action_type", "module", "detail", "status", "client_info",
            "request_path", "request_method", "response_code", "duration_ms",
            "sensitive",
        ]
        parts = []
        for field in fields:
            val = entry.get(field)
            if val is None:
                val = ""
            elif isinstance(val, (dict, list)):
                val = json.dumps(val, ensure_ascii=False, sort_keys=True)
            else:
                val = str(val)
            parts.append(val)
        concat = "|".join(parts)
        return hashlib.sha256(concat.encode("utf-8")).hexdigest()

    async def query(self, query: AuditLogQuery) -> AuditLogPageResult:
        clauses: list[str] = []
        params: list[Any] = []
        fts_clause = None
        fts_params: list[Any] = []

        if query.user_id is not None:
            clauses.append("user_id LIKE ?")
            params.append(f"%{query.user_id}%")

        if query.username is not None:
            clauses.append("username LIKE ?")
            params.append(f"%{query.username}%")

        if query.time_start is not None:
            clauses.append("timestamp >= ?")
            params.append(query.time_start.isoformat())

        if query.time_end is not None:
            clauses.append("timestamp <= ?")
            params.append(query.time_end.isoformat())

        if query.action_types is not None and len(query.action_types) > 0:
            placeholders = ",".join("?" for _ in query.action_types)
            clauses.append(f"action_type IN ({placeholders})")
            params.extend([at.value for at in query.action_types])

        if query.modules is not None and len(query.modules) > 0:
            placeholders = ",".join("?" for _ in query.modules)
            clauses.append(f"module IN ({placeholders})")
            params.extend([m.value for m in query.modules])

        if query.status is not None:
            clauses.append("status = ?")
            params.append(query.status)

        if query.request_path is not None:
            clauses.append("request_path LIKE ?")
            params.append(f"%{query.request_path}%")

        if query.keyword is not None and query.keyword.strip():
            fts_query = query.keyword
            if not re.search(r'\b(AND|OR|NOT)\b', fts_query, re.IGNORECASE):
                fts_query = " AND ".join(re.findall(r'\S+', fts_query))
            fts_clause = fts_query
            fts_params = [fts_query]

        where_parts = []
        all_params: list[Any] = []

        if fts_clause is not None:
            base_from = """
                audit_logs_fts fts
                JOIN audit_logs al ON al.rowid = fts.rowid
            """
            where_parts.append("fts.audit_logs_fts MATCH ?")
            all_params.extend(fts_params)
        else:
            base_from = "audit_logs al"

        if clauses:
            where_parts.extend(clauses)
            all_params.extend(params)

        where_sql = ""
        if where_parts:
            where_sql = "WHERE " + " AND ".join(where_parts)

        count_sql = f"SELECT COUNT(*) FROM {base_from} {where_sql}"
        async with self._db.execute(count_sql, all_params) as cursor:
            row = await cursor.fetchone()
            total = row[0]

        order_field = query.order_by if query.order_by in self.VALID_ORDER_FIELDS else "timestamp"
        order_dir = "ASC" if query.order_dir.lower() == "asc" else "DESC"

        offset = (query.page - 1) * query.page_size
        data_sql = f"""
            SELECT al.id, al.user_id, al.username, al.timestamp, al.ip_address,
                   al.action_type, al.module, al.detail, al.status, al.client_info,
                   al.request_path, al.request_method, al.response_code, al.duration_ms,
                   al.hash, al.sensitive
            FROM {base_from}
            {where_sql}
            ORDER BY al.{order_field} {order_dir}
            LIMIT ? OFFSET ?
        """
        data_params = all_params + [query.page_size, offset]

        items = []
        async with self._db.execute(data_sql, data_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                items.append(self._row_to_audit_log(r))

        total_pages = ceil(total / query.page_size) if total > 0 else 0

        return AuditLogPageResult(
            total=total,
            page=query.page,
            page_size=query.page_size,
            total_pages=total_pages,
            items=items,
        )

    async def get_by_id(self, log_id: str) -> Optional[AuditLog]:
        async with self._db.execute(
            """
            SELECT id, user_id, username, timestamp, ip_address,
                   action_type, module, detail, status, client_info,
                   request_path, request_method, response_code, duration_ms,
                   hash, sensitive
            FROM audit_logs WHERE id = ?
            """,
            (log_id,),
        ) as cursor:
            row = await cursor.fetchone()
            if row is None:
                return None
            return self._row_to_audit_log(row)

    async def verify_integrity(self, log_id: str = None) -> AuditLogVerifyResult:
        if log_id is not None:
            async with self._db.execute(
                "SELECT id, user_id, username, timestamp, ip_address, action_type, module, detail, status, client_info, request_path, request_method, response_code, duration_ms, hash, sensitive FROM audit_logs WHERE id = ?",
                (log_id,),
            ) as cursor:
                row = await cursor.fetchone()
                if row is None:
                    return AuditLogVerifyResult(total=0, passed=0, failed=0, failed_ids=[])
                entry_dict = self._row_to_entry_dict(row)
                stored_hash = entry_dict.pop("hash", "")
                computed = self.compute_hash(entry_dict)
                if computed == stored_hash:
                    return AuditLogVerifyResult(total=1, passed=1, failed=0, failed_ids=[])
                else:
                    return AuditLogVerifyResult(total=1, passed=0, failed=1, failed_ids=[log_id])

        total = 0
        passed = 0
        failed = 0
        failed_ids: list[str] = []

        async with self._db.execute(
            "SELECT id, user_id, username, timestamp, ip_address, action_type, module, detail, status, client_info, request_path, request_method, response_code, duration_ms, hash, sensitive FROM audit_logs"
        ) as cursor:
            rows = await cursor.fetchall()
            for row in rows:
                total += 1
                entry_dict = self._row_to_entry_dict(row)
                stored_hash = entry_dict.pop("hash", "")
                computed = self.compute_hash(entry_dict)
                if computed == stored_hash:
                    passed += 1
                else:
                    failed += 1
                    failed_ids.append(row["id"])

        return AuditLogVerifyResult(
            total=total,
            passed=passed,
            failed=failed,
            failed_ids=failed_ids,
        )

    async def get_statistics(
        self,
        time_start: datetime = None,
        time_end: datetime = None,
        granularity: str = "hour",
    ) -> AuditLogStatistics:
        granularity_map = {
            "minute": "%Y-%m-%d %H:%M",
            "hour": "%Y-%m-%d %H:00",
            "day": "%Y-%m-%d",
            "week": "%Y-W%W",
            "month": "%Y-%m",
        }
        fmt = granularity_map.get(granularity, "%Y-%m-%d %H:00")

        time_clauses = []
        time_params: list[Any] = []
        if time_start is not None:
            time_clauses.append("timestamp >= ?")
            time_params.append(time_start.isoformat())
        if time_end is not None:
            time_clauses.append("timestamp <= ?")
            time_params.append(time_end.isoformat())

        where_sql = ""
        if time_clauses:
            where_sql = "WHERE " + " AND ".join(time_clauses)

        trend_sql = f"""
            SELECT strftime('{fmt}', timestamp) AS bucket, COUNT(*) AS count
            FROM audit_logs
            {where_sql}
            GROUP BY bucket
            ORDER BY bucket
        """
        trend = []
        async with self._db.execute(trend_sql, time_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                trend.append({"bucket": r[0], "count": r[1]})

        module_sql = f"""
            SELECT module, COUNT(*) AS count
            FROM audit_logs
            {where_sql}
            GROUP BY module
            ORDER BY count DESC
        """
        module_distribution = []
        async with self._db.execute(module_sql, time_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                module_distribution.append({"module": r[0], "count": r[1]})

        action_sql = f"""
            SELECT action_type, COUNT(*) AS count
            FROM audit_logs
            {where_sql}
            GROUP BY action_type
            ORDER BY count DESC
        """
        action_distribution = []
        async with self._db.execute(action_sql, time_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                action_distribution.append({"action_type": r[0], "count": r[1]})

        user_sql = f"""
            SELECT user_id, username, COUNT(*) AS count
            FROM audit_logs
            {where_sql}
            GROUP BY user_id, username
            ORDER BY count DESC
            LIMIT 10
        """
        user_ranking = []
        async with self._db.execute(user_sql, time_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                user_ranking.append({
                    "user_id": r[0],
                    "username": r[1],
                    "count": r[2],
                })

        anomaly_sql = f"""
            SELECT user_id, username, COUNT(*) AS fail_count,
                   MIN(timestamp) AS first_fail, MAX(timestamp) AS last_fail
            FROM audit_logs
            WHERE status = 'failure' {"AND " + " AND ".join(time_clauses) if time_clauses else ""}
            GROUP BY user_id, username
            HAVING fail_count > 10
               AND strftime('%s', MAX(timestamp)) - strftime('%s', MIN(timestamp)) <= 300
            ORDER BY fail_count DESC
        """
        anomalies = []
        async with self._db.execute(anomaly_sql, time_params) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                anomalies.append({
                    "user_id": r[0],
                    "username": r[1],
                    "fail_count": r[2],
                    "first_fail": r[3],
                    "last_fail": r[4],
                })

        return AuditLogStatistics(
            trend=trend,
            module_distribution=module_distribution,
            action_distribution=action_distribution,
            user_ranking=user_ranking,
            anomalies=anomalies,
        )

    async def export_excel(self, query: AuditLogQuery, output_path: str) -> str:
        from openpyxl import Workbook

        result = await self.query(query)

        exports_dir = Path(output_path).parent
        exports_dir.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "审计日志"

        headers = [
            "ID", "用户ID", "用户名", "时间戳", "IP地址", "操作类型",
            "模块", "详情", "状态", "客户端信息", "请求路径", "请求方法",
            "响应码", "耗时(ms)", "哈希", "敏感操作",
        ]
        ws.append(headers)

        for item in result.items:
            detail_str = ""
            if item.detail is not None:
                detail_str = json.dumps(item.detail, ensure_ascii=False)
            ws.append([
                item.id, item.user_id, item.username, item.timestamp.isoformat() if item.timestamp else "",
                item.ip_address, item.action_type.value if item.action_type else "",
                item.module.value if item.module else "", detail_str,
                item.status, item.client_info, item.request_path, item.request_method,
                item.response_code, item.duration_ms, item.hash,
                "是" if item.sensitive else "否",
            ])

        meta_ws = wb.create_sheet(title="元数据")
        meta_ws.append(["导出时间", datetime.now().isoformat()])
        meta_ws.append(["总记录数", str(result.total)])
        meta_ws.append(["页码", str(result.page)])
        meta_ws.append(["每页条数", str(result.page_size)])
        if query.user_id:
            meta_ws.append(["用户ID", query.user_id])
        if query.username:
            meta_ws.append(["用户名", query.username])
        if query.action_types:
            meta_ws.append(["操作类型", ",".join(at.value for at in query.action_types)])
        if query.modules:
            meta_ws.append(["模块", ",".join(m.value for m in query.modules)])
        if query.status:
            meta_ws.append(["状态", query.status])

        wb.save(output_path)
        return output_path

    async def export_csv(self, query: AuditLogQuery, output_path: str) -> str:
        result = await self.query(query)

        exports_dir = Path(output_path).parent
        exports_dir.mkdir(parents=True, exist_ok=True)

        with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
            f.write(f"# 导出时间: {datetime.now().isoformat()}\n")
            f.write(f"# 总记录数: {result.total}\n")
            f.write(f"# 查询条件: user_id={query.user_id or ''}, username={query.username or ''}, status={query.status or ''}\n")

            writer = csv.writer(f)
            headers = [
                "ID", "用户ID", "用户名", "时间戳", "IP地址", "操作类型",
                "模块", "详情", "状态", "客户端信息", "请求路径", "请求方法",
                "响应码", "耗时(ms)", "哈希", "敏感操作",
            ]
            writer.writerow(headers)

            for item in result.items:
                detail_str = ""
                if item.detail is not None:
                    detail_str = json.dumps(item.detail, ensure_ascii=False)
                writer.writerow([
                    item.id, item.user_id, item.username,
                    item.timestamp.isoformat() if item.timestamp else "",
                    item.ip_address, item.action_type.value if item.action_type else "",
                    item.module.value if item.module else "", detail_str,
                    item.status, item.client_info, item.request_path, item.request_method,
                    item.response_code, item.duration_ms, item.hash,
                    "是" if item.sensitive else "否",
                ])

        return output_path

    async def export_pdf(self, query: AuditLogQuery, output_path: str) -> str:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        result = await self.query(query)

        exports_dir = Path(output_path).parent
        exports_dir.mkdir(parents=True, exist_ok=True)

        doc = SimpleDocTemplate(
            output_path,
            pagesize=landscape(A4),
            leftMargin=20, rightMargin=20,
            topMargin=40, bottomMargin=40,
        )

        styles = getSampleStyleSheet()
        elements = []

        header_text = f"审计日志导出 - 导出时间: {datetime.now().isoformat()} - 总记录数: {result.total}"
        elements.append(Paragraph(header_text, styles["Normal"]))
        elements.append(Spacer(1, 12))

        table_headers = [
            "用户ID", "用户名", "时间戳", "IP", "操作类型",
            "模块", "状态", "请求路径", "响应码", "耗时(ms)",
        ]
        table_data = [table_headers]

        for item in result.items:
            row = [
                item.user_id,
                item.username,
                item.timestamp.isoformat() if item.timestamp else "",
                item.ip_address,
                item.action_type.value if item.action_type else "",
                item.module.value if item.module else "",
                item.status,
                item.request_path,
                str(item.response_code),
                str(item.duration_ms),
            ]
            table_data.append(row)

        col_widths = [60, 60, 120, 80, 50, 60, 40, 120, 40, 50]
        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("TOPPADDING", (0, 0), (-1, 0), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))

        elements.append(table)

        def add_page_number(canvas, doc):
            page_num = canvas.getPageNumber()
            text = f"第 {page_num} 页"
            canvas.drawRightString(A4[0] - 20, 20, text)

        doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)
        return output_path

    async def rotate_and_archive(self) -> Optional[str]:
        cutoff = datetime.now() - timedelta(days=self.ARCHIVE_RETENTION_YEARS * 365)

        async with self._db.execute(
            "SELECT COUNT(*) FROM audit_logs WHERE timestamp < ?",
            (cutoff.isoformat(),),
        ) as cursor:
            row = await cursor.fetchone()
            count = row[0]

        if count == 0:
            return None

        async with self._db.execute(
            "SELECT MIN(timestamp), MAX(timestamp) FROM audit_logs WHERE timestamp < ?",
            (cutoff.isoformat(),),
        ) as cursor:
            row = await cursor.fetchone()
            min_ts = row[0]
            max_ts = row[1]

        if min_ts is None:
            return None

        period = min_ts[:7].replace("-", "")
        project_root = Path(__file__).resolve().parent.parent.parent.parent
        archive_dir = project_root / "data" / "audit_archive"
        archive_dir.mkdir(parents=True, exist_ok=True)
        archive_path = archive_dir / f"audit_archive_{period}.db"

        archive_db = await aiosqlite.connect(str(archive_path))
        await archive_db.executescript("""
            CREATE TABLE IF NOT EXISTS audit_logs (
                id TEXT PRIMARY KEY,
                user_id TEXT NOT NULL DEFAULT 'anonymous',
                username TEXT NOT NULL DEFAULT 'anonymous',
                timestamp TEXT NOT NULL,
                ip_address TEXT NOT NULL DEFAULT '',
                action_type TEXT NOT NULL,
                module TEXT NOT NULL,
                detail TEXT,
                status TEXT NOT NULL DEFAULT 'success',
                client_info TEXT NOT NULL DEFAULT '',
                request_path TEXT NOT NULL DEFAULT '',
                request_method TEXT NOT NULL DEFAULT '',
                response_code INTEGER NOT NULL DEFAULT 200,
                duration_ms REAL NOT NULL DEFAULT 0.0,
                hash TEXT NOT NULL DEFAULT '',
                sensitive INTEGER NOT NULL DEFAULT 0
            );
        """)
        await archive_db.commit()

        async with self._db.execute(
            """
            SELECT id, user_id, username, timestamp, ip_address, action_type, module,
                   detail, status, client_info, request_path, request_method,
                   response_code, duration_ms, hash, sensitive
            FROM audit_logs WHERE timestamp < ?
            """,
            (cutoff.isoformat(),),
        ) as cursor:
            rows = await cursor.fetchall()
            for r in rows:
                await archive_db.execute(
                    """
                    INSERT OR IGNORE INTO audit_logs (id, user_id, username, timestamp, ip_address,
                        action_type, module, detail, status, client_info,
                        request_path, request_method, response_code, duration_ms, hash, sensitive)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    tuple(r),
                )
        await archive_db.commit()
        await archive_db.close()

        await self._db.execute(
            "DELETE FROM audit_logs WHERE timestamp < ?",
            (cutoff.isoformat(),),
        )
        await self._db.commit()

        return str(archive_path)

    async def backup(self) -> str:
        project_root = Path(__file__).resolve().parent.parent.parent.parent
        backup_dir = project_root / "data" / "audit_backup"
        backup_dir.mkdir(parents=True, exist_ok=True)

        now = datetime.now()
        backup_name = f"audit_backup_{now.strftime('%Y%m%d_%H%M%S')}.db"
        backup_path = backup_dir / backup_name

        if self._db is not None:
            await self._db.close()
            self._db = None

        shutil.copy2(self._db_path, str(backup_path))

        self._db = await aiosqlite.connect(self._db_path)
        self._db.row_factory = aiosqlite.Row

        cutoff = now - timedelta(days=self.BACKUP_RETENTION_DAYS)
        for f in backup_dir.iterdir():
            if f.is_file() and f.name.startswith("audit_backup_") and f.suffix == ".db":
                try:
                    date_str = f.stem.replace("audit_backup_", "")
                    file_time = datetime.strptime(date_str, "%Y%m%d_%H%M%S")
                    if file_time < cutoff:
                        f.unlink()
                except (ValueError, OSError):
                    pass

        return str(backup_path)

    async def list_archives(self) -> list[AuditArchiveInfo]:
        project_root = Path(__file__).resolve().parent.parent.parent.parent
        archive_dir = project_root / "data" / "audit_archive"
        if not archive_dir.exists():
            return []

        archives = []
        for f in sorted(archive_dir.glob("audit_archive_*.db")):
            if not f.is_file():
                continue
            record_count = 0
            period_start = ""
            period_end = ""
            try:
                async with aiosqlite.connect(str(f)) as adb:
                    adb.row_factory = aiosqlite.Row
                    try:
                        async with adb.execute("SELECT COUNT(*) FROM audit_logs") as cursor:
                            row = await cursor.fetchone()
                            record_count = row[0]
                        async with adb.execute(
                            "SELECT MIN(timestamp), MAX(timestamp) FROM audit_logs"
                        ) as cursor:
                            row = await cursor.fetchone()
                            if row[0] is not None:
                                period_start = row[0]
                            if row[1] is not None:
                                period_end = row[1]
                    except Exception:
                        pass
            except Exception:
                pass

            archives.append(AuditArchiveInfo(
                filename=f.name,
                size_bytes=f.stat().st_size,
                record_count=record_count,
                period_start=period_start,
                period_end=period_end,
            ))

        return archives

    async def get_recent_logs(self, limit: int = 5) -> list[AuditLog]:
        async with self._db.execute(
            """
            SELECT id, user_id, username, timestamp, ip_address,
                   action_type, module, detail, status, client_info,
                   request_path, request_method, response_code, duration_ms,
                   hash, sensitive
            FROM audit_logs
            ORDER BY timestamp DESC
            LIMIT ?
            """,
            (limit,),
        ) as cursor:
            rows = await cursor.fetchall()
            return [self._row_to_audit_log(r) for r in rows]

    def _row_to_audit_log(self, row: aiosqlite.Row) -> AuditLog:
        detail = row["detail"]
        if detail is not None:
            try:
                detail = json.loads(detail)
            except (json.JSONDecodeError, TypeError):
                pass

        timestamp_str = row["timestamp"]
        timestamp_val = None
        if timestamp_str:
            try:
                timestamp_val = datetime.fromisoformat(timestamp_str)
            except (ValueError, TypeError):
                timestamp_val = datetime.now()

        sensitive_val = bool(row["sensitive"])

        return AuditLog(
            id=row["id"],
            user_id=row["user_id"],
            username=row["username"],
            timestamp=timestamp_val,
            ip_address=row["ip_address"],
            action_type=row["action_type"],
            module=row["module"],
            detail=detail,
            status=row["status"],
            client_info=row["client_info"],
            request_path=row["request_path"],
            request_method=row["request_method"],
            response_code=row["response_code"],
            duration_ms=row["duration_ms"],
            hash=row["hash"],
            sensitive=sensitive_val,
        )

    def _row_to_entry_dict(self, row: aiosqlite.Row) -> dict:
        detail = row["detail"]
        if detail is not None:
            try:
                detail = json.loads(detail)
            except (json.JSONDecodeError, TypeError):
                pass

        return {
            "id": row["id"],
            "user_id": row["user_id"],
            "username": row["username"],
            "timestamp": row["timestamp"],
            "ip_address": row["ip_address"],
            "action_type": row["action_type"],
            "module": row["module"],
            "detail": detail,
            "status": row["status"],
            "client_info": row["client_info"],
            "request_path": row["request_path"],
            "request_method": row["request_method"],
            "response_code": row["response_code"],
            "duration_ms": row["duration_ms"],
            "hash": row["hash"],
            "sensitive": bool(row["sensitive"]),
        }


audit_log_service = AuditLogService()
